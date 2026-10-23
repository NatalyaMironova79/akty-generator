import streamlit as st
import openpyxl
from datetime import datetime, timedelta
from docxtpl import DocxTemplate
from docx import Document
from docx.shared import Pt
from copy import deepcopy
from num2words import num2words
import zipfile
import io
import re
import os

# --- Склонение ФИО ---
try:
    import pymorphy3
    _morph = pymorphy3.MorphAnalyzer()
except Exception:
    _morph = None

ORG_FORMS = ["ЧТУП", "ОДО", "ООО", "ОАО", "ЗАО", "УП", "ИП", "ПТ", "СП", "АО"]

ORG_FORMS_ABLT = {
    "ЧТУП": "Частным торговым унитарным предприятием",
    "УП": "Унитарным предприятием",
    "ООО": "Обществом с ограниченной ответственностью",
    "ОДО": "Обществом с дополнительной ответственностью",
    "ЗАО": "Закрытым акционерным обществом",
    "ОАО": "Открытым акционерным обществом",
    "АО": "Акционерным обществом",
    "ПТ": "Полным товариществом",
    "СП": "Совместным предприятием",
    "ИП": "Индивидуальным предпринимателем",
}


def parse_executor(executor_name_raw):
    if not executor_name_raw:
        return ("", "")
    s = str(executor_name_raw).strip()
    for form in ORG_FORMS:
        pattern = rf'(?:^|\s)["«]?{form}["»]?\s*$'
        if re.search(pattern, s, flags=re.IGNORECASE):
            clean = re.sub(pattern, '', s, flags=re.IGNORECASE).strip()
            return (form.upper(), clean)
    return ("", s)


def _detect_gender(fio):
    parts = str(fio).split()
    if len(parts) >= 3:
        otch = parts[2].lower()
        if otch.endswith(("вна", "ична", "инична")):
            return "f"
        if otch.endswith(("вич", "ич")):
            return "m"
    if len(parts) >= 2:
        name = parts[1].lower()
        female_names = {"анна","мария","елена","ольга","татьяна","наталья","ирина",
                        "светлана","екатерина","юлия","марина","галина","людмила",
                        "валентина","лариса","наталия","вера","надежда","любовь",
                        "алла","инна","оксана","жанна","софья","софия","дарья",
                        "полина","алина","кристина","виктория","анастасия"}
        if name in female_names:
            return "f"
    return None


def _decline_fio(fio, case):
    if not fio:
        return ""
    if _morph is None:
        return fio
    parts = str(fio).split()
    if len(parts) < 2:
        return fio
    declined = []
    for i, word in enumerate(parts):
        if not word or not word[0].isalpha():
            declined.append(word)
            continue
        if i == 0:
            wl = word.lower()
            if not wl.endswith(("ов","ев","ёв","ин","ын","ский","цкий",
                                "ой","ая","ий","ова","ева","ёва","ина","ына","а","я")):
                declined.append(word)
                continue
        parsed = _morph.parse(word)[0]
        inflected = parsed.inflect({case})
        if inflected:
            result = inflected.word
            if word[0].isupper():
                result = result.capitalize()
            declined.append(result)
        else:
            declined.append(word)
    return " ".join(declined)


def decline_fio_ablt(fio):
    return _decline_fio(fio, "ablt")


def decline_fio_gent(fio):
    return _decline_fio(fio, "gent")


def decline_position(position):
    if not position:
        return ""
    s = str(position).strip()
    if not s:
        return ""
    wl = s.lower()
    if wl.endswith(("а", "я", "ого", "его")):
        return s
    if _morph is None:
        return s
    parsed = _morph.parse(s)[0]
    inflected = parsed.inflect({"gent"})
    if inflected:
        result = inflected.word
        if s[0].isupper():
            result = result.capitalize()
        return result
    return s


MONTHS_RU = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
             "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]


def month_year_ru(date_str):
    try:
        dt = datetime.strptime(date_str, "%d.%m.%Y")
        return f"{MONTHS_RU[dt.month - 1]} {dt.year}"
    except (ValueError, TypeError):
        return ""


def short_name(full_name):
    if not full_name:
        return ""
    s = str(full_name).strip()
    prefixes = [
        "Индивидуальный предприниматель", "ИП",
        "Унитарное предприятие", "УП",
        "Общество с ограниченной ответственностью", "ООО", "ОДО", "ЗАО", "ОАО",
    ]
    for pref in prefixes:
        if s.lower().startswith(pref.lower()):
            s = s[len(pref):].strip()
            s = s.lstrip('«»"\' -—')
            break
    parts = s.split()
    if len(parts) >= 3:
        return f"{parts[0]} {parts[1][0]}.{parts[2][0]}."
    if len(parts) == 2:
        return f"{parts[0]} {parts[1][0]}."
    return s


# --- Streamlit UI ---
st.set_page_config(page_title="Генератор актов", layout="centered")
st.title("Генератор актов сдачи-приёмки")
st.write("Загрузите XLS с данными — получите ZIP с готовыми актами.")

# Шаблон docx — зашит в приложение
TEMPLATE_PATH = "акт_шаблон.docx"

if not os.path.exists(TEMPLATE_PATH):
    st.error(f"Не найден шаблон {TEMPLATE_PATH}. Обратитесь к администратору.")
    st.stop()

xls_file = st.file_uploader("XLS с данными", type=["xlsx"])


def fmt_date(value):
    if value is None or value == "": return ""
    if isinstance(value, datetime): return value.strftime("%d.%m.%Y")
    if isinstance(value, (int, float)):
        base = datetime(1899, 12, 30)
        return (base + timedelta(days=int(value))).strftime("%d.%m.%Y")
    s = str(value).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d.%m.%Y"):
        try: return datetime.strptime(s, fmt).strftime("%d.%m.%Y")
        except ValueError: pass
    return s


def fmt_money(value):
    if value is None or value == "": return "0,00"
    try: num = float(value)
    except (ValueError, TypeError): return str(value)
    return f"{num:,.2f}".replace(",", " ").replace(".", ",")


def money_to_words(value):
    try: num = float(value)
    except (ValueError, TypeError): return ""
    rub = int(num); kop = int(round((num - rub) * 100))
    rub_words = num2words(rub, lang="ru")
    last_two = rub % 100; last_one = rub % 10
    if last_two in (11,12,13,14): rub_word = "рублей"
    elif last_one == 1: rub_word = "рубль"
    elif last_one in (2,3,4): rub_word = "рубля"
    else: rub_word = "рублей"
    if kop % 10 == 1 and kop % 100 != 11: kop_word = "копейка"
    elif kop % 10 in (2,3,4) and kop % 100 not in (12,13,14): kop_word = "копейки"
    else: kop_word = "копеек"
    return f"{rub_words} {rub_word} {kop:02d} {kop_word}"


def normalize_vat_mode(value):
    s = str(value).strip().lower()
    return "с НДС" if ("с ндс" in s and "без" not in s) else "без НДС"


def get_vat_note(vat_mode, vat_rate, total_with_vat, executor_type=""):
    if vat_mode == "без НДС":
        if executor_type == "ИП":
            return "без НДС согласно ч.2, ст. 112 НК РБ"
        else:
            return "без НДС согласно пп.1.2, п.1, ст. 326 НК РБ"
    try:
        rate = float(vat_rate)
    except (ValueError, TypeError):
        rate = 20.0
    vat_amount = total_with_vat * rate / (100 + rate)
    return f"в том числе НДС {fmt_money(vat_amount)} ({money_to_words(vat_amount)}) руб."


def set_cell(cell, text):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    run.font.size = Pt(10)


def clone_row_after(table, row_idx):
    src_tr = table.rows[row_idx]._tr
    new_tr = deepcopy(src_tr)
    src_tr.addnext(new_tr)
    return table.rows[row_idx + 1]


def remove_row(table, row_idx):
    table._tbl.remove(table.rows[row_idx]._tr)


# --- Основная логика ---
if xls_file:
    if st.button("Сгенерировать акты", type="primary"):
        with st.spinner("Генерируем..."):
            with open("_uploaded.xlsx", "wb") as f:
                f.write(xls_file.getbuffer())

            wb = openpyxl.load_workbook("_uploaded.xlsx", data_only=True)

            def read_sheet(name):
                ws = wb[name]
                headers = [c.value for c in ws[1]]
                rows = []
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if row[0] is None: continue
                    rows.append(dict(zip(headers, row[:len(headers)])))
                return rows

            try:
                partners = read_sheet("partner")
                serv_rows = read_sheet("serv")
                car_rows = read_sheet("cars")
            except Exception as e:
                st.error(f"Ошибка в XLS: {e}")
                st.stop()

            zip_buffer = io.BytesIO()
            file_count = 0

            with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
                for partner in partners:
                    org_id = partner["OrgID"]
                    executor_name_raw = str(partner.get("ExecutorName") or "").strip()

                    parsed_type, clean_name = parse_executor(executor_name_raw)
                    executor_type = str(partner.get("ExecutorType") or "").strip()
                    if not executor_type:
                        executor_type = parsed_type
                    executor_name = clean_name

                    executor_rep = partner.get("ExecutorRepresentative") or ""
                    executor_pos = str(partner.get("ExecutorPosition") or "").strip()
                    executor_basis = partner.get("ExecutorBasis") or ""

                    if not executor_pos:
                        executor_pos = "" if executor_type == "ИП" else "директор"

                    name_instr_manual = str(partner.get("ExecutorNameInstr") or "").strip()
                    rep_gen_manual = str(partner.get("ExecutorRepresentativeGen") or "").strip()

                    org_serv = [r for r in serv_rows if r["OrgID"] == org_id]
                    org_cars = [r for r in car_rows if r["OrgID"] == org_id]

                    if executor_type == "ИП":
                        name_instr = name_instr_manual or decline_fio_ablt(executor_name)
                        gender = _detect_gender(executor_name)
                        participle = "действующей" if gender == "f" else "действующим"
                        basis_clean = str(executor_basis).strip()
                        if not basis_clean:
                            basis_clean = f"Свидетельства о государственной регистрации № {org_id}"
                        preamble = (f"Индивидуальным предпринимателем {name_instr}, "
                                    f"{participle} на основании {basis_clean}")
                        executor_rep_short = short_name(executor_name)
                    else:
                        full_type_ablt = ORG_FORMS_ABLT.get(executor_type, executor_type)
                        pos_gen = decline_position(executor_pos)
                        rep_gen = rep_gen_manual or decline_fio_gent(executor_rep)
                        if rep_gen:
                            preamble = (f"{full_type_ablt} {executor_name}, в лице {pos_gen} {rep_gen}, "
                                        f"действующего на основании {executor_basis}")
                        else:
                            preamble = (f"{full_type_ablt} {executor_name}, в лице {pos_gen}, "
                                        f"действующего на основании {executor_basis}")
                        executor_rep_short = short_name(executor_rep)

                    serv_total = 0.0
                    serv_data = []
                    for idx, r in enumerate(org_serv, start=1):
                        cnt = int(r["CarsCount"] or 0)
                        price = float(r["PricePerCar"] or 0)
                        total = cnt * price
                        serv_total += total
                        serv_data.append({
                            "n": idx, "cars_count": cnt,
                            "period_from": fmt_date(r["PeriodFrom"]),
                            "period_to": fmt_date(r["PeriodTo"]),
                            "price_per_car": fmt_money(price),
                            "total_cost": fmt_money(total),
                        })

                    report_month = month_year_ru(fmt_date(partner["ReportPeriodFrom"]))
                    cars_total = 0.0
                    cars_data = []
                    for idx, r in enumerate(org_cars, start=1):
                        reward = float(r["RewardAmount"] or 0)
                        cars_total += reward
                        cars_data.append({
                            "n": idx, "car_reg_number": r["CarRegNumber"],
                            "report_period": report_month,
                            "trips_count": int(r["TripsCount"] or 0),
                            "reward_amount": fmt_money(reward),
                        })

                    vat_mode = normalize_vat_mode(partner["VATMode"])
                    vat_note_serv = get_vat_note(vat_mode, partner["VATRate"], serv_total, executor_type)
                    vat_note_cars = get_vat_note(vat_mode, partner["VATRate"], cars_total, executor_type)
                    report_to = fmt_date(partner["ReportPeriodTo"])

                    ctx = {
                        "appendix_number": partner["AppendixNumber"] or 1,
                        "contract_number": partner["ContractNumber"],
                        "appendix_contract_date": fmt_date(partner["AppendixContractDate"] or partner["ContractDate"]),
                        "act_date": report_to,
                        "preamble": preamble,
                        "executor_representative_short": executor_rep_short,
                        "contract_date": fmt_date(partner["ContractDate"]),
                        "total_serv_cost": fmt_money(serv_total),
                        "total_serv_cost_words": money_to_words(serv_total),
                        "total_reward_amount": fmt_money(cars_total),
                        "total_reward_amount_words": money_to_words(cars_total),
                        "vat_note_serv": vat_note_serv,
                        "vat_note_cars": vat_note_cars,
                    }

                    doc = DocxTemplate(TEMPLATE_PATH)
                    doc.render(ctx)
                    doc.save("_temp.docx")

                    d = Document("_temp.docx")

                    t1 = d.tables[0]
                    remove_row(t1, 3)
                    remove_row(t1, 2)
                    for row in serv_data:
                        new_row = clone_row_after(t1, 1)
                        set_cell(new_row.cells[0], row["n"])
                        set_cell(new_row.cells[1], row["cars_count"])
                        set_cell(new_row.cells[2], row["period_from"])
                        set_cell(new_row.cells[3], row["period_to"])
                        set_cell(new_row.cells[4], row["price_per_car"])
                        set_cell(new_row.cells[5], row["total_cost"])
                    set_cell(t1.rows[-1].cells[5], fmt_money(serv_total))

                    t2 = d.tables[1]
                    remove_row(t2, 2)
                    remove_row(t2, 1)
                    for row in cars_data:
                        new_row = clone_row_after(t2, 0)
                        set_cell(new_row.cells[0], row["n"])
                        set_cell(new_row.cells[1], row["car_reg_number"])
                        set_cell(new_row.cells[2], row["report_period"])
                        set_cell(new_row.cells[3], row["trips_count"])
                        set_cell(new_row.cells[4], row["reward_amount"])
                    set_cell(t2.rows[-1].cells[4], fmt_money(cars_total))

                    safe_name = re.sub(r"[^\w\-]", "_", str(executor_name))
                    filename = f"Акт_{safe_name}_{report_to.replace('.', '-')}.docx"
                    d.save("_result.docx")
                    zf.write("_result.docx", filename)
                    file_count += 1

            st.success(f"Готово! Сгенерировано актов: {file_count}")
            st.download_button(
                label="Скачать ZIP с актами",
                data=zip_buffer.getvalue(),
                file_name="акты.zip",
                mime="application/zip",
                type="primary",
            )
